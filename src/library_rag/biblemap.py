"""The Bible map: biblical places, people and events, and where each event happened.

The data is a hand-curated set of spreadsheets kept outside this repo (places
and people exported from Notion, events from Logos with IDs resolved against
both). This module turns those files into the `biblemap` schema and reads it
back for the page.

Three stages, the same split as bible.py and for the same reason -- each is
understandable and testable on its own:

    read_sources()   files -> Dataset          (no database)
    validate()       Dataset -> list of problems
    load()           Dataset -> Postgres       (no filesystem)

Alongside the parsed tables, every source file is also kept VERBATIM -- every
row, every column, as text -- in biblemap.source_files / source_rows. That copy
is what the "Show in original data" viewer reads: when the map looks wrong, the
first question is whether the spreadsheet says so, and answering it should not
mean opening Excel and hunting for a row.
"""
import csv
import json
from dataclasses import dataclass, field
from pathlib import Path

# The file names inside the source directory. Globs for the two Notion exports,
# because Notion appends a hash to every export ("Places Data 17ed4e..._all.csv")
# and re-exporting changes it.
PLACES_GLOB = "Places-Data/Places Data *_all.csv"
NEW_PLACES_FILE = "Places-Data/New Places (added during resolution).csv"
PEOPLE_GLOB = "People Data/People Data *_all.csv"
EVENTS_FILE = "Bible Events with Resolved PersonIDs and PlaceIDs.xlsx"

# The columns the map actually reads, per source file. The viewer marks these,
# so it is obvious which of the 43 People columns can change what the map shows.
USED_COLUMNS = {
    "events": ["Passage (Logos Data)", "Title", "Description", "Icon",
               "PlaceID", "RoutePlaceID", "resolved_people_ids"],
    "places": ["PlaceID", "PlaceName", "AltName", "AltSpelling", "Lat", "Lng",
               "Comments", "Verses"],
    "new_places": ["PlaceID", "PlaceName", "Lat", "Lng", "Comments"],
    "people": ["PersonID", "Name", "AltName", "Descriptor", "SubjectType",
               "Gender", "Verses"],
}

# Columns whose cells are lists of IDs pointing into another file -- the
# foreign keys between the spreadsheets. The viewer turns each ID into a link to
# the row it names. "Duplicate IDs" is not read by the map, but it is the list of
# candidate people the resolution chose between, which is exactly what someone
# checking a resolution wants to click through.
LINK_COLUMNS = {
    "events": {"PlaceID": "place", "RoutePlaceID": "place",
               "resolved_people_ids": "person", "Duplicate IDs": "person"},
}

# Which source files a map record can come from, for locate().
RECORD_SOURCES = {
    "event": ["events"],
    "place": ["places", "new_places"],
    "person": ["people"],
}


@dataclass
class Source:
    """One source file, verbatim.

    `rows` are (row_num, record_id, cells). row_num is the row as a spreadsheet
    program numbers it -- the header is row 1 -- so "row 41" in the viewer is
    row 41 when the file is opened in Excel. record_id is the event, place or
    person the row became, or None for a row that became nothing (a blank row,
    a place with no ID).
    """
    key: str
    filename: str
    columns: list
    rows: list = field(default_factory=list)

    def records(self):
        for row_num, _, cells in self.rows:
            yield row_num, dict(zip(self.columns, cells))


@dataclass
class Dataset:
    # Row tuples in COPY column order -- see load().
    places: list = field(default_factory=list)
    people: list = field(default_factory=list)
    events: list = field(default_factory=list)
    event_places: list = field(default_factory=list)
    event_route: list = field(default_factory=list)
    event_people: list = field(default_factory=list)
    sources: list = field(default_factory=list)
    # Source rows dropped because they had no ID -- the places export carries
    # one ("malt"). Counted rather than silently skipped, so the CLI can say so.
    skipped: dict = field(default_factory=dict)


# ------------------------------------------------------------------ parsing --

def _cell(value) -> str:
    """A spreadsheet value as the text a person would see in the cell.

    openpyxl hands back numbers as numbers, and an ID typed into Excel can
    arrive as 354.0; the viewer should show 354.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _int(value):
    """A cell as an int ID, or None."""
    text = str(value or "").strip()
    return int(text) if text.isdigit() else None


def _ids(value) -> list:
    """"354,951, 534" -> [354, 951, 534]. Raises on anything that is not an ID,
    so a leftover "TODO: ..." note cannot load as a silently shorter list."""
    out = []
    for part in str(value or "").split(","):
        part = part.strip()
        if not part:
            continue
        if not part.isdigit():
            raise ValueError(f"not an ID: {part!r}")
        out.append(int(part))
    return out


def _coord(lat, lng):
    """(lat, lng) as floats, or (None, None) when the source has no real point.

    The source writes "unknown" three ways -- blank, 0/0, and a guess with a
    question mark ("37?") -- and all three have to become NULL, because the map
    would otherwise plot them.
    """
    try:
        la, ln = float(str(lat).strip()), float(str(lng).strip())
    except ValueError:
        return None, None
    if la == 0 and ln == 0:
        return None, None
    if not (-90 <= la <= 90 and -180 <= ln <= 180):
        return None, None
    return la, ln


def _text(value) -> str:
    """A cell as clean text. The Notion exports write missing values as NULL."""
    text = str(value or "").strip()
    return "" if text == "NULL" else text


def _one(source_dir: Path, pattern: str) -> Path:
    matches = sorted(source_dir.glob(pattern))
    if len(matches) != 1:
        raise FileNotFoundError(
            f"expected exactly one {pattern!r} in {source_dir}, found {len(matches)}"
        )
    return matches[0]


def read_csv_source(key: str, path: Path, id_column: str) -> Source:
    # utf-8-sig: both Notion exports begin with a byte-order mark, which would
    # otherwise become part of the first column's name ("﻿PlaceID").
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        columns = next(reader)
        source = Source(key, path.name, columns)
        id_at = columns.index(id_column)
        for row_num, cells in enumerate(reader, start=2):
            cells = cells + [""] * (len(columns) - len(cells))
            source.rows.append((row_num, _int(cells[id_at]), cells))
    return source


def read_places(sources) -> tuple:
    rows, skipped = [], 0
    for source in sources:
        for _, r in source.records():
            place_id = _int(r.get("PlaceID"))
            if place_id is None:
                skipped += 1
                continue
            lat, lng = _coord(r.get("Lat", ""), r.get("Lng", ""))
            alt = ", ".join(
                _text(r.get(k)) for k in ("AltName", "AltSpelling") if _text(r.get(k))
            )
            rows.append((
                place_id, _text(r.get("PlaceName")), alt, lat, lng,
                _text(r.get("Comments")), _text(r.get("Verses")),
            ))
    return rows, skipped


def read_people(source: Source) -> tuple:
    rows, skipped = [], 0
    for _, r in source.records():
        person_id = _int(r.get("PersonID"))
        if person_id is None:
            skipped += 1
            continue
        rows.append((
            person_id, _text(r.get("Name")), _text(r.get("AltName")),
            _text(r.get("Descriptor")), _text(r.get("SubjectType")),
            _text(r.get("Gender")), _text(r.get("Verses")),
        ))
    return rows, skipped


def read_events_source(path: Path) -> Source:
    """The events workbook, verbatim: every row as the text in its cells.

    record_id is left unset here and assigned by parse_events, which is the one
    place that decides what an event's id is.

    openpyxl is imported here rather than at the top so the web app, which never
    reads a spreadsheet, does not need it importable.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb.active.iter_rows(values_only=True)
        columns = [_cell(h) for h in next(sheet)]
        raw = [[_cell(v) for v in row] for row in sheet]
    finally:
        wb.close()

    source = Source("events", path.name, columns)
    for row_num, cells in enumerate(raw, start=2):
        source.rows.append((row_num, None, (cells + [""] * len(columns))[:len(columns)]))
    return source


def is_event_row(columns, cells) -> bool:
    """Whether a row of the events file is an event at all. A row with neither a
    passage nor a title is a blank separator, not an event with things missing."""
    r = dict(zip(columns, cells))
    return bool(_text(r.get("Passage (Logos Data)")) or _text(r.get("Title")))


def parse_events(source: Source, ds: Dataset) -> None:
    """The events rows -> ds.events and the three link tables.

    Each row keeps whatever id it already carries (record_id), so an event's id
    survives rows being inserted above it; a row that has none -- every row on
    first import, and any row added since -- takes the next free number. `seq`
    is the row's position among event rows, and that is what the timeline reads.

    Writes the assigned ids back into source.rows, so the caller can store them.
    """
    taken = {rid for _, rid, _ in source.rows if rid is not None}
    next_id = max(taken, default=0) + 1
    assigned, seq = [], 0
    for row_num, record_id, cells in source.rows:
        if not is_event_row(source.columns, cells):
            assigned.append((row_num, None, cells))
            continue
        if record_id is None:
            record_id, next_id = next_id, next_id + 1
        assigned.append((row_num, record_id, cells))
        seq += 1
        r = dict(zip(source.columns, cells))
        ds.events.append((
            record_id, seq, _text(r.get("Passage (Logos Data)")), _text(r.get("Title")),
            _text(r.get("Description")), _text(r.get("Icon")),
        ))
        try:
            places = _ids(r.get("PlaceID"))
            route = _ids(r.get("RoutePlaceID"))
            people = _ids(r.get("resolved_people_ids"))
        except ValueError as err:
            passage = _text(r.get("Passage (Logos Data)")) or _text(r.get("Title"))
            raise ValueError(f"events row {row_num} ({passage}): {err}") from None
        # dict.fromkeys: de-duplicate, keeping first-seen order.
        for pos, pid in enumerate(dict.fromkeys(places)):
            ds.event_places.append((record_id, pid, pos))
        for pos, pid in enumerate(route):
            ds.event_route.append((record_id, pos, pid))
        for pos, pid in enumerate(dict.fromkeys(people)):
            ds.event_people.append((record_id, pid, pos))
    source.rows = assigned


def dataset_from_sources(sources) -> Dataset:
    """Source files -> the tables the map is built from. No files, no database.

    The one parse path: the loader reads spreadsheets into Sources and calls
    this, and an edit reads Sources back out of Postgres and calls this. Editing
    a cell therefore does exactly what editing the spreadsheet would have done.
    """
    by_key = {s.key: s for s in sources}
    ds = Dataset(sources=list(sources))
    places = [by_key[k] for k in ("places", "new_places") if k in by_key]
    ds.places, ds.skipped["places"] = read_places(places)
    ds.people, ds.skipped["people"] = read_people(by_key["people"])
    parse_events(by_key["events"], ds)
    return ds


def read_sources(source_dir) -> Dataset:
    """Every source file in `source_dir` -> one Dataset. Touches no database."""
    source_dir = Path(source_dir)
    sources = [read_events_source(source_dir / EVENTS_FILE),
               read_csv_source("places", _one(source_dir, PLACES_GLOB), "PlaceID")]
    if (source_dir / NEW_PLACES_FILE).exists():
        sources.append(
            read_csv_source("new_places", source_dir / NEW_PLACES_FILE, "PlaceID"))
    sources.append(read_csv_source("people", _one(source_dir, PEOPLE_GLOB), "PersonID"))
    return dataset_from_sources(sources)


# --------------------------------------------------------------- validation --

def validate(ds: Dataset) -> list:
    """Every reason not to load `ds`, as readable lines. [] means load it.

    Collected up front rather than left to the foreign keys: Postgres stops at
    the first violation and names a constraint, where this names every bad ID
    and the event it is in, which is what someone fixing the spreadsheet needs.
    """
    problems = []
    for label, rows in (("place", ds.places), ("person", ds.people)):
        seen, dupes = set(), set()
        for r in rows:
            (dupes if r[0] in seen else seen).add(r[0])
        if dupes:
            problems.append(f"duplicate {label} IDs: {sorted(dupes)[:20]}")

    place_ids = {r[0] for r in ds.places}
    person_ids = {r[0] for r in ds.people}
    passage = {r[0]: r[2] for r in ds.events}
    checks = (
        ("place", ds.event_places, 1, place_ids),
        ("route place", ds.event_route, 2, place_ids),
        ("person", ds.event_people, 1, person_ids),
    )
    for label, links, idx, known in checks:
        for link in links:
            if link[idx] not in known:
                problems.append(
                    f"event {link[0]} ({passage[link[0]]}): unknown {label} ID {link[idx]}"
                )
    if not ds.events:
        problems.append("no events found")
    return problems


# --------------------------------------------------------------------- load --

class Invalid(Exception):
    """A change that would leave the map inconsistent. Carries every problem."""

    def __init__(self, problems):
        super().__init__("; ".join(problems[:5]))
        self.problems = problems


def _copy(conn, table: str, columns: str, rows) -> None:
    with conn.cursor() as cur:
        with cur.copy(f"COPY biblemap.{table} ({columns}) FROM STDIN") as copy:
            for row in rows:
                copy.write_row(row)


DERIVED = (
    ("places", "place_id, name, alt_names, lat, lng, comments, verses"),
    ("people", "person_id, name, alt_name, descriptor, subject_type, gender, verses"),
    ("events", "event_id, seq, passage, title, description, icon"),
    ("event_places", "event_id, place_id, position"),
    ("event_route", "event_id, position, place_id"),
    ("event_people", "event_id, person_id, position"),
)


def load_derived(conn, ds: Dataset) -> dict:
    """Replace the map's tables with what `ds` parsed to. Does not commit.

    TRUNCATE-and-COPY rather than a diff: the parse is a pure function of the
    source rows, so rewriting the result is trivially correct where reconciling
    it against whatever was there before is a second implementation to get wrong.
    """
    conn.execute(
        "TRUNCATE biblemap.event_places, biblemap.event_route, biblemap.event_people, "
        "biblemap.events, biblemap.people, biblemap.places"
    )
    written = {}
    for table, columns in DERIVED:
        rows = getattr(ds, table)
        _copy(conn, table, columns, rows)
        written[table] = len(rows)
    return written


def load_sources(conn, ds: Dataset) -> dict:
    """Replace the verbatim copy of the source files. Does not commit."""
    conn.execute("TRUNCATE biblemap.source_rows, biblemap.source_files")
    _copy(conn, "source_files", "file_key, filename, columns, row_count",
          [(s.key, s.filename, json.dumps(s.columns), len(s.rows)) for s in ds.sources])
    rows = [(s.key, row_num, record_id, json.dumps(cells))
            for s in ds.sources for row_num, record_id, cells in s.rows]
    _copy(conn, "source_rows", "file_key, row_num, record_id, cells", rows)
    return {"source_files": len(ds.sources), "source_rows": len(rows)}


def load(conn, ds: Dataset) -> dict:
    """Replace everything in the biblemap schema with `ds`. Returns row counts.

    Same shape as bible.load, for the same reasons: TRUNCATE-and-reload because
    a freshly read set of files is the whole truth, COPY because the database
    may be across the network, one transaction so a failure leaves the old map
    intact. The verbatim source copy is replaced in the same transaction, so the
    viewer can never show a file other than the one the map was built from.
    """
    written = {**load_derived(conn, ds), **load_sources(conn, ds)}
    conn.commit()
    return written


# ------------------------------------------------------------------ editing --
# The source rows are the truth and the map is derived from them, so every edit
# is the same three steps: change a row, re-parse every row, replace the derived
# tables -- in ONE transaction, so an edit that would break the map leaves
# nothing behind. See migrations/0020_biblemap_editing.sql.


def sources_from_db(conn) -> list:
    """The stored source files, as the same Source objects the readers produce."""
    files = conn.execute(
        "SELECT file_key, filename, columns FROM biblemap.source_files"
    ).fetchall()
    sources = []
    for key, filename, columns in files:
        rows = conn.execute(
            "SELECT row_num, record_id, cells FROM biblemap.source_rows "
            "WHERE file_key = %s ORDER BY row_num",
            (key,),
        ).fetchall()
        sources.append(Source(key, filename, columns, [tuple(r) for r in rows]))
    return sources


def rebuild(conn) -> dict:
    """Re-derive the map from the stored source rows. Does not commit.

    Raises Invalid, having written nothing, when the rows do not make a
    consistent map -- an event pointing at a place that is not there, a
    duplicated ID. The caller rolls back.
    """
    sources = sources_from_db(conn)
    ds = dataset_from_sources(sources)
    problems = validate(ds)
    if problems:
        raise Invalid(problems)
    written = load_derived(conn, ds)
    # parse_events may have given new rows their ids; store them, or the next
    # rebuild would hand the same event a different id.
    events = next(s for s in ds.sources if s.key == "events")
    for row_num, record_id, _ in events.rows:
        conn.execute(
            "UPDATE biblemap.source_rows SET record_id = %s "
            "WHERE file_key = 'events' AND row_num = %s AND record_id IS DISTINCT FROM %s",
            (record_id, row_num, record_id),
        )
    return written


def _row(conn, key: str, row_num: int):
    row = conn.execute(
        "SELECT row_num, record_id, cells FROM biblemap.source_rows "
        "WHERE file_key = %s AND row_num = %s",
        (key, row_num),
    ).fetchone()
    if row is None:
        raise KeyError(f"{key} has no row {row_num}")
    return row


def _columns(conn, key: str) -> list:
    row = conn.execute(
        "SELECT columns FROM biblemap.source_files WHERE file_key = %s", (key,)
    ).fetchone()
    if row is None:
        raise KeyError(f"no source file {key!r}")
    return row[0]


def _log(conn, key, row_num, action, editor, column=None, old=None, new=None,
         snapshot=None):
    conn.execute(
        "INSERT INTO biblemap.source_edits "
        "(editor, file_key, row_num, action, column_name, old_value, new_value, "
        " row_snapshot) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
        (editor, key, row_num, action, column, old, new,
         json.dumps(snapshot) if snapshot is not None else None),
    )


def _touch(conn, key, row_num, editor):
    conn.execute(
        "UPDATE biblemap.source_rows SET edited_at = now(), editor = %s "
        "WHERE file_key = %s AND row_num = %s",
        (editor, key, row_num),
    )


def update_row(conn, key: str, row_num: int, changes: dict, editor=None) -> dict:
    """Change cells in one row, then re-derive the map. Commits, or rolls back.

    `changes` is {column: text}; a column the file does not have is an error
    rather than a silently ignored key.
    """
    try:
        columns = _columns(conn, key)
        _, _, cells = _row(conn, key, row_num)
        cells = list(cells)
        for column, value in changes.items():
            if column not in columns:
                raise KeyError(f"{key} has no column {column!r}")
            i = columns.index(column)
            old, cells[i] = cells[i], str(value)
            if old != cells[i]:
                _log(conn, key, row_num, "update", editor, column, old, cells[i])
        conn.execute(
            "UPDATE biblemap.source_rows SET cells = %s WHERE file_key = %s AND row_num = %s",
            (json.dumps(cells), key, row_num),
        )
        _touch(conn, key, row_num, editor)
        written = rebuild(conn)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"row": row_num, "cells": cells, "written": written}


def next_id(conn, key: str) -> int:
    """The next free ID for a new place or person row.

    Places and people added in the app continue the 10000+ range the resolution
    file already uses, so an ID says at a glance whether it came from the
    original data.
    """
    column = "PersonID" if key == "people" else "PlaceID"
    keys = ("people",) if key == "people" else ("places", "new_places")
    used = {10000}
    for k in keys:
        try:
            columns = _columns(conn, k)
        except KeyError:
            continue  # new_places is optional
        if column not in columns:
            continue
        at = columns.index(column)
        for (cells,) in conn.execute(
            "SELECT cells FROM biblemap.source_rows WHERE file_key = %s", (k,)
        ).fetchall():
            value = _int(cells[at]) if at < len(cells) else None
            if value is not None:
                used.add(value)
    return max(used) + 1


def insert_row(conn, key: str, after_row: int, editor=None) -> dict:
    """Add a blank row directly after `after_row`, then re-derive the map.

    Row numbers stay the spreadsheet's: everything below shifts down by one, in
    one statement, which is why 0020 made the primary key deferrable.
    """
    try:
        columns = _columns(conn, key)
        _row(conn, key, after_row)  # refuse to insert after a row that is not there
        conn.execute("SET CONSTRAINTS biblemap.source_rows_pkey DEFERRED")
        conn.execute(
            "UPDATE biblemap.source_rows SET row_num = row_num + 1 "
            "WHERE file_key = %s AND row_num > %s",
            (key, after_row),
        )
        cells = [""] * len(columns)
        # A new place or person is nothing without an ID, and asking someone to
        # invent a free one is asking them to check 1,285 rows first.
        if key in ("places", "new_places", "people"):
            id_column = "PersonID" if key == "people" else "PlaceID"
            if id_column in columns:
                cells[columns.index(id_column)] = str(next_id(conn, key))
        row_num = after_row + 1
        conn.execute(
            "INSERT INTO biblemap.source_rows "
            "(file_key, row_num, record_id, cells, origin, edited_at, editor) "
            "VALUES (%s, %s, NULL, %s, 'app', now(), %s)",
            (key, row_num, json.dumps(cells), editor),
        )
        conn.execute(
            "UPDATE biblemap.source_files SET row_count = row_count + 1 WHERE file_key = %s",
            (key,),
        )
        _log(conn, key, row_num, "insert", editor, snapshot=cells)
        # A blank events row parses as a separator rather than an event, and a
        # blank place row has an ID and no name -- both are consistent, so the
        # map can be rebuilt now and the row filled in afterwards.
        written = rebuild(conn)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"row": row_num, "cells": cells, "written": written}


def delete_row(conn, key: str, row_num: int, editor=None) -> dict:
    try:
        _, _, cells = _row(conn, key, row_num)
        conn.execute(
            "DELETE FROM biblemap.source_rows WHERE file_key = %s AND row_num = %s",
            (key, row_num),
        )
        conn.execute("SET CONSTRAINTS biblemap.source_rows_pkey DEFERRED")
        conn.execute(
            "UPDATE biblemap.source_rows SET row_num = row_num - 1 "
            "WHERE file_key = %s AND row_num > %s",
            (key, row_num),
        )
        conn.execute(
            "UPDATE biblemap.source_files SET row_count = row_count - 1 WHERE file_key = %s",
            (key,),
        )
        _log(conn, key, row_num, "delete", editor, snapshot=list(cells))
        written = rebuild(conn)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"row": row_num, "written": written}


def edit_count(conn) -> int:
    """How many stored rows the app has changed -- what makes a re-import from
    the spreadsheets destructive."""
    return conn.execute(
        "SELECT count(*) FROM biblemap.source_rows "
        "WHERE origin = 'app' OR edited_at IS NOT NULL"
    ).fetchone()[0]


def export(conn, out_dir) -> list:
    """Write the stored rows back out as spreadsheets. Returns the paths.

    The round trip that keeps biblemap_project meaningful once the app holds the
    truth: export, replace the files, and a fresh --load reproduces exactly what
    the app has.
    """
    import openpyxl

    out_dir = Path(out_dir)
    written = []
    for key, filename, columns in conn.execute(
        "SELECT file_key, filename, columns FROM biblemap.source_files ORDER BY file_key"
    ).fetchall():
        rows = conn.execute(
            "SELECT cells FROM biblemap.source_rows WHERE file_key = %s ORDER BY row_num",
            (key,),
        ).fetchall()
        # Back into the layout read_sources expects, so an export can be dropped
        # straight over biblemap_project and re-loaded.
        folder = {"places": "Places-Data", "new_places": "Places-Data",
                  "people": "People Data"}.get(key, "")
        path = out_dir / folder / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.suffix == ".xlsx":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(columns)
            for (cells,) in rows:
                ws.append(cells)
            wb.save(path)
        else:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(columns)
                w.writerows(cells for (cells,) in rows)
        written.append(path)
    return written


# ------------------------------------------------------------------ reading --

def loaded(conn) -> bool:
    """Whether the schema exists AND holds events. See bible.loaded for why
    to_regclass rather than just querying."""
    if conn.execute("SELECT to_regclass('biblemap.events')").fetchone()[0] is None:
        return False
    return conn.execute("SELECT EXISTS (SELECT 1 FROM biblemap.events)").fetchone()[0]


def counts(conn) -> dict:
    row = conn.execute(
        """
        SELECT (SELECT count(*) FROM biblemap.events),
               (SELECT count(*) FROM biblemap.people),
               (SELECT count(*) FROM biblemap.places),
               (SELECT count(*) FROM biblemap.places WHERE lat IS NULL)
        """
    ).fetchone()
    return dict(zip(("events", "people", "places", "places_without_coords"), row))


def data(conn) -> dict:
    """Everything the page needs, in one payload.

    One call rather than an endpoint per view: the three views are filters over
    the same ~1,700 events, and filtering in the browser makes switching views
    instant. Places and people are limited to the ones some event mentions --
    the rest have nothing to show.
    """
    events = conn.execute(
        """
        SELECT e.event_id, e.passage, e.title, e.description, e.icon,
               COALESCE((SELECT array_agg(place_id ORDER BY position)
                         FROM biblemap.event_places p WHERE p.event_id = e.event_id), '{}'),
               COALESCE((SELECT array_agg(place_id ORDER BY position)
                         FROM biblemap.event_route r WHERE r.event_id = e.event_id), '{}'),
               COALESCE((SELECT array_agg(person_id ORDER BY position)
                         FROM biblemap.event_people p WHERE p.event_id = e.event_id), '{}')
        FROM biblemap.events e
        ORDER BY e.seq
        """
    ).fetchall()
    places = conn.execute(
        """
        SELECT place_id, name, lat, lng
        FROM biblemap.places
        WHERE place_id IN (SELECT place_id FROM biblemap.event_places
                           UNION SELECT place_id FROM biblemap.event_route)
        ORDER BY name
        """
    ).fetchall()
    people = conn.execute(
        """
        SELECT person_id, name, descriptor
        FROM biblemap.people
        WHERE person_id IN (SELECT person_id FROM biblemap.event_people)
        ORDER BY name
        """
    ).fetchall()
    return {
        "events": [
            {"id": i, "passage": passage, "title": title, "description": desc,
             "icon": icon, "places": pl, "route": rt, "people": pp}
            for i, passage, title, desc, icon, pl, rt, pp in events
        ],
        "places": [
            {"id": i, "name": name, "lat": lat, "lng": lng}
            for i, name, lat, lng in places
        ],
        "people": [
            {"id": i, "name": name, "descriptor": descriptor}
            for i, name, descriptor in people
        ],
    }


# ----------------------------------------------------------- source viewer --

def source_files(conn) -> list:
    """The loaded source files, in load order: events, places, people."""
    return conn.execute(
        """
        SELECT file_key, filename, columns, row_count, loaded_at
        FROM biblemap.source_files
        ORDER BY array_position(ARRAY['events','places','new_places','people'], file_key)
        """
    ).fetchall()


def record_names(conn) -> dict:
    """{"place": {id: name}, "person": {id: name}} for every loaded record, so
    the viewer can label an ID link with who or what it points at."""
    return {
        "place": dict(conn.execute("SELECT place_id, name FROM biblemap.places").fetchall()),
        "person": dict(conn.execute("SELECT person_id, name FROM biblemap.people").fetchall()),
    }


def locate(conn, kind: str, record_id: int):
    """(file_key, row_num) of the source row a map record came from, or None."""
    keys = RECORD_SOURCES.get(kind)
    if not keys:
        return None
    return conn.execute(
        """
        SELECT file_key, row_num FROM biblemap.source_rows
        WHERE file_key = ANY(%s) AND record_id = %s
        ORDER BY array_position(%s, file_key)
        LIMIT 1
        """,
        (keys, record_id, keys),
    ).fetchone()


def source_page(conn, key: str, offset: int, limit: int, q: str = "") -> tuple:
    """One window of a source file's rows, optionally filtered by `q`.

    Returns (total_matching, rows) with rows as (row_num, record_id, cells).
    The filter is a literal, case-insensitive match against any cell -- the same
    escaping as bible.search, for the same reason.
    """
    where, params = "file_key = %s", [key]
    if q:
        pattern = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        where += " AND cells::text ILIKE %s ESCAPE '\\'"
        params.append(f"%{pattern}%")
    total = conn.execute(
        f"SELECT count(*) FROM biblemap.source_rows WHERE {where}", params
    ).fetchone()[0]
    rows = conn.execute(
        f"""
        SELECT row_num, record_id, cells, origin, edited_at, editor
        FROM biblemap.source_rows
        WHERE {where} ORDER BY row_num OFFSET %s LIMIT %s
        """,
        [*params, offset, limit],
    ).fetchall()
    return total, rows


def source_row_index(conn, key: str, row_num: int) -> int:
    """How many rows of `key` come before `row_num` -- the offset that puts it
    on screen."""
    return conn.execute(
        "SELECT count(*) FROM biblemap.source_rows WHERE file_key = %s AND row_num < %s",
        (key, row_num),
    ).fetchone()[0]
